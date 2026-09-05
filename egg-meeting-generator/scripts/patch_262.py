# -*- coding: utf-8 -*-
"""EGG第262次会议：按角色接龙修订 Agenda Excel + PPT背景板（幂等，可重跑）

前置修复：原始 262 背景板 pptx 携带 3 个 261 遗留的孤儿 slide 部件
（slide17 备稿演讲4 / slide21 备稿个评3 / slide22 备稿个评4），
仅在 presentation.xml.rels 中挂链、未进入 sldIdLst。python-pptx 保存时会
把它们重新序列化并与其他页的 partname 冲突，导致页面上出现 261 残留页。
因此先做 zip 级清理，再交给 python-pptx 编辑。
"""
import sys, re, copy as cp, zipfile
from pptx import Presentation
from pptx.oxml.ns import qn
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

BASE = 'F:/workbuddy/2026-08-28-22-56-26/output/'
SRC_PPTX = BASE + '_orig.pptx'
XLSX = BASE + 'EGG262次会议Agenda.xlsx'
PPTX = BASE + 'EGG - 262次会议背景板.pptx'
CLEAN = BASE + '_clean.pptx'

# ============================================================
# 接龙数据（第262次会议 / 2026-08-30 周日 9:30-12:00）
# ============================================================
GRAMMARIAN = '向莹'          # 语法官
PHOTOGRAPHER = '待定'        # 拍照官（接龙为 [烟花]）
MANAGER = 'Kim@EGG'          # 会议经理

# ============================================================
# Step 0: 清理孤儿 slide 部件
# ============================================================
def orphan_slide_parts(path):
    z = zipfile.ZipFile(path)
    pres = z.read('ppt/presentation.xml').decode('utf-8')
    rels = z.read('ppt/_rels/presentation.xml.rels').decode('utf-8')
    rmap = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
    used = {rmap[i] for i in re.findall(r'<p:sldId[^>]*r:id="(rId\d+)"', pres)}
    allp = [n for n in z.namelist()
            if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)]
    orphans = sorted(n for n in allp if n[len('ppt/'):] not in used)
    z.close()
    return orphans


ORPHANS = orphan_slide_parts(SRC_PPTX)
print('孤儿 slide 部件:', ORPHANS or '无')
if ORPHANS:
    drop = set(ORPHANS)
    for n in ORPHANS:
        drop.add(n.replace('slides/', 'slides/_rels/') + '.rels')
    src = zipfile.ZipFile(SRC_PPTX)
    with zipfile.ZipFile(CLEAN, 'w', zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            if item.filename in drop:
                continue
            if item.filename == '[Content_Types].xml':
                xml = src.read(item.filename).decode('utf-8')
                for n in ORPHANS:
                    xml = re.sub(
                        r'<Override PartName="/' + re.escape(n) + r'"[^>]*/>', '', xml)
            elif item.filename == 'ppt/_rels/presentation.xml.rels':
                xml = src.read(item.filename).decode('utf-8')
                for n in ORPHANS:
                    tgt = n[len('ppt/'):]
                    xml = re.sub(r'<Relationship [^>]*Target="' + re.escape(tgt) + r'"[^>]*/>', '', xml)
            else:
                out.writestr(item, src.read(item.filename))
                continue
            out.writestr(item.filename, xml.encode('utf-8'))
    src.close()
    print('已清理 →', CLEAN)
else:
    import shutil
    shutil.copy2(SRC_PPTX, CLEAN)

# ============================================================
# Part 1: Excel
# ============================================================
wb = openpyxl.load_workbook(XLSX)
ws = wb['正面']

ws['I13'] = GRAMMARIAN      # 语法官介绍
ws['I26'] = GRAMMARIAN      # 语法官报告

ws['B19'] = (
    '口号：EGG！破壳，诞生无限！\n\n'
    '愿景：打造中国最具企业家精神的演讲实验室。让每一次开口，都诞生无限可能。\n\n'
    '使命：提供积极互助互益的学习体验，通过赋能会员提高演讲力和领导力，获得更大自信及个人成长\n\n'
    f'会议经理：{MANAGER}\n'
    'IT场控：\n'
    '接待官：待定\n'
    f'拍照官：{PHOTOGRAPHER}'
)
wb.save(XLSX)
print('Excel OK')

# ============================================================
# Part 2: PPT
# ============================================================
prs = Presentation(CLEAN)


def set_para_text(para, new_text):
    """含 \\x0b(软换行) 时按 a:br 重建，避免 _x000B_ 字面量"""
    p = para._p
    end = p.find(qn('a:endParaRPr'))
    rPr = None
    for r in p.findall(qn('a:r')):
        el = r.find(qn('a:rPr'))
        if el is not None:
            rPr = cp.deepcopy(el)
            break
    for child in list(p):
        if child.tag in (qn('a:r'), qn('a:br'), qn('a:fld')):
            p.remove(child)

    def add(el):
        if end is not None:
            end.addprevious(el)
        else:
            p.append(el)

    for i, seg in enumerate(new_text.split('\x0b')):
        if i > 0:
            br = p.makeelement(qn('a:br'), {})
            if rPr is not None:
                br.append(cp.deepcopy(rPr))
            add(br)
        r = p.makeelement(qn('a:r'), {})
        if rPr is not None:
            r.append(cp.deepcopy(rPr))
        t = r.makeelement(qn('a:t'), {})
        t.text = seg
        r.append(t)
        add(r)


def replace_in_shape(shape, old, new):
    if not shape.has_text_frame:
        return False
    hit = False
    for para in shape.text_frame.paragraphs:
        if old in para.text:
            set_para_text(para, para.text.replace(old, new))
            hit = True
    return hit


def replace_all(slide, reps):
    for shape in slide.shapes:
        for old, new in reps:
            replace_in_shape(shape, old, new)


def find_slide(prs, title):
    for i, s in enumerate(prs.slides):
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() == title:
                return i
    return None


# --- 1) 删除「提问官提问」页（本次接龙无提问官，Excel 议程亦无此环节）---
i_ask = find_slide(prs, '提问官提问')
if i_ask is not None:
    lst = prs.slides._sldIdLst
    el = list(lst)[i_ask]
    prs.part.drop_rel(el.get(qn('r:id')))
    lst.remove(el)
    print('已删除 提问官提问 页')
else:
    print('提问官提问 页不存在（可能已删除）')

# --- 2) 总评团队页：重建为 4 行（去掉提问官，语法官=向莹）---
TEAM = ['总评官：无忌哥', '时间官: 徐朝霞', f'语法官：{GRAMMARIAN}', '哼哈官：Chris']
done = False
for shape in prs.slides[find_slide(prs, '总评团队介绍')].shapes:
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            if '总评官' in para.text:
                set_para_text(para, '\x0b'.join(TEAM))
                done = True
print('总评团队页重建:', done)

# --- 3) 语法官报告页 ---
replace_all(prs.slides[find_slide(prs, '语法官报告')], [('待定', GRAMMARIAN)])

# --- 4) EGG十分钟：补上分享主题（复用大咖分享页副标题样式）---
s10 = prs.slides[find_slide(prs, 'EGG十分钟')]
s_share = prs.slides[find_slide(prs, '大咖分享')]
sub_src, share_name = None, None
for sh in s_share.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip().startswith('《'):
        sub_src = sh
for sh in s_share.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == '程雪@EGG':
        share_name = sh
name_ph = None
for sh in s10.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == '金灿灿':
        name_ph = sh
if sub_src is None or name_ph is None or share_name is None:
    print('!! EGG十分钟 主题添加失败')
elif any(sh.has_text_frame and '品项羽' in sh.text_frame.text for sh in s10.shapes):
    print('EGG十分钟 主题已存在')
else:
    new_sub = cp.deepcopy(sub_src._element)
    for p in new_sub.iter(qn('a:p')):
        for r in p.findall(qn('a:r')):
            t = r.find(qn('a:t'))
            if t is not None:
                t.text = '品项羽-不寻常的饭局'
                break
    name_ph._element.addnext(new_sub)
    # 姓名下移到与大咖分享页一致的位置，避免与副标题重叠
    name_ph.top, name_ph.left = share_name.top, share_name.left
    print('EGG十分钟 主题已添加')

prs.save(PPTX)
print('PPT OK')

# ============================================================
# 验证
# ============================================================
z = zipfile.ZipFile(PPTX)
bad = [n for n in z.namelist()
       if n.startswith('ppt/slides/slide') and n.endswith('.xml') and b'_x000B_' in z.read(n)]
print('含_x000B_字面量的页:', bad if bad else '无')

prs2 = Presentation(PPTX)
print('残留孤儿页:', orphan_slide_parts(PPTX) or '无')
print('\n=== PPT 共', len(prs2.slides), '页 ===')
for i, s in enumerate(prs2.slides):
    t = ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    print(f'{i+1:2d}. {t.replace(chr(11), "|").replace(chr(10), "|")[:70]}')
