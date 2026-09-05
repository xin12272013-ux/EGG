import openpyxl
import shutil
from copy import copy
from datetime import time

# Copy template
src = 'F:/workbuddy/2026-07-28-14-26-36/EGG257_template_url.xlsx'
dst = 'F:/workbuddy/2026-07-28-14-26-36/EGG258次会议Agenda.xlsx'
shutil.copy2(src, dst)

wb = openpyxl.load_workbook(dst)

# ==================== 正面 Sheet ====================
ws = wb['正面']

# --- Header ---
ws['C3'] = '第258次中文线下会议【2026年8月2日】'
ws['C4'] = '会议主题：《世界500强扭亏为盈的秘密》'

# --- Sidebar info (B5) ---
ws['B5'] = (
    '会议时间：\n 2026年8月2日（周日）\n上午 9:30 - 中午12:00\n\n'
    '会议地点：\n上海徐汇区云锦路181号龙华街道水岸党群服务中心214室（龙华地铁站附近）\n\n'
    '费用：所有来宾和角色全免费\n\n'
    '注意事项： \n◆ 会议4大禁忌：\n 政治 * 宗教 * 性 * 低俗话题 \n◆ 会场纪律:\n'
    '- 请将手机静音；\n- 不大声喧哗和随意走动；\n- 上下舞台请与主持人握手；\n- 为发言者鼓掌。'
)

# --- Set start time ---
ws['C7'] = time(9, 30)

# --- Role/Agenda rows ---
# Row 7: 暖场游戏
ws['E7'] = '暖场游戏'
ws['F7'] = 3; ws['G7'] = 4; ws['H7'] = 5
ws['I7'] = '待定'

# Row 8: 主持人开场
ws['E8'] = '会议主持人开场'
ws['F8'] = 1; ws['G8'] = 1.5; ws['H8'] = 2
ws['I8'] = 'Kim@EGG'

# Row 9: 时间官介绍
ws['E9'] = '时间官介绍'
ws['F9'] = 1; ws['G9'] = 1.5; ws['H9'] = 2
ws['I9'] = 'Rachel'

# Row 10: 提问官介绍
ws['E10'] = '提问官介绍'
ws['F10'] = 1; ws['G10'] = 1.5; ws['H10'] = 2
ws['I10'] = '吴琼'

# Row 11: 主题分享
ws['E11'] = '《世界500强扭亏为盈的秘密》'
ws['F11'] = 35; ws['G11'] = 40; ws['H11'] = 45
ws['I11'] = '慕艺'

# Row 12: 提问官提问 (I12 formula =I10 stays)
ws['E12'] = '提问官提问'
ws['F12'] = 1; ws['G12'] = 2; ws['H12'] = 3

# Row 13: 来宾介绍
ws['E13'] = '来宾介绍'
ws['F13'] = 5; ws['G13'] = 6; ws['H13'] = 7
ws['I13'] = '待定@EGG'

# Row 14: 中场休息 (keep as-is)
ws['E14'] = '中场休息'
ws['F14'] = 4; ws['G14'] = 5; ws['H14'] = 6

# Row 15: 即兴演讲
ws['B15'] = '每日一词:'
ws['E15'] = '即兴演讲'
ws['F15'] = 13; ws['G15'] = 14; ws['H15'] = 15
ws['I15'] = 'Amy@EGG'

# Row 16: 备稿1
ws['E16'] = '备稿1'
ws['F16'] = 5; ws['G16'] = 6; ws['H16'] = 7
ws['I16'] = '程雪@EGG'

# Row 17: 备稿2
ws['E17'] = '备稿2'
ws['F17'] = 5; ws['G17'] = 6; ws['H17'] = 7
ws['I17'] = '快乐哥@EGG'

# Row 18: 备稿3 (was EGG十分钟)
ws['E18'] = '备稿3'
ws['F18'] = 5; ws['G18'] = 6; ws['H18'] = 7
ws['I18'] = 'Kim@EGG'

# Row 19: 即兴评估
ws['E19'] = '即兴评估'
ws['F19'] = 4; ws['G19'] = 5; ws['H19'] = 6
ws['I19'] = '陶然'

# Row 20: 备稿评估1
ws['E20'] = '备稿评估1'
ws['F20'] = 2; ws['G20'] = 2.5; ws['H20'] = 3
ws['I20'] = 'Mark'

# Row 21: 备稿评估2
ws['E21'] = '备稿评估2'
ws['F21'] = 2; ws['G21'] = 2.5; ws['H21'] = 3
ws['I21'] = '韩羽良'

# Row 22: 备稿评估3 (was 时间官报告)
ws['E22'] = '备稿评估3'
ws['F22'] = 2; ws['G22'] = 2.5; ws['H22'] = 3
ws['I22'] = '待定'

# Row 23: 时间官报告 (was 投票&来宾反馈)
ws['E23'] = '时间官报告'
ws['F23'] = 1; ws['G23'] = 1.5; ws['H23'] = 2
ws['I23'] = 'Rachel'

# Row 24: 投票&来宾反馈 (was 主席致辞)
ws['E24'] = '投票&来宾反馈'
ws['F24'] = 4; ws['G24'] = 5; ws['H24'] = 6
ws['I24'] = 'Ocean'

# Row 25: 主席致辞&入会规则介绍 (was 通告&颁奖)
ws['E25'] = '主席致辞&入会规则介绍'
ws['F25'] = 4; ws['G25'] = 5; ws['H25'] = 6
ws['I25'] = '快乐哥@EGG'

# Row 26: 通告&颁奖 (new - was previously empty or different content)
ws['E26'] = '通告&颁奖'
ws['F26'] = 2; ws['G26'] = 3; ws['H26'] = 4
ws['I26'] = '快乐哥@EGG'

# --- Update B16 sidebar ---
ws['B16'] = (
    '口号：EGG！破壳，诞生无限！\n\n'
    '愿景：打造中国最具企业家精神的演讲实验室。让每一次开口，都诞生无限可能。\n\n'
    '使命：提供积极互助互益的学习体验，通过赋能会员提高演讲力和领导力，获得更大自信及个人成长\n\n'
    '会议经理：Kim\n'
    'IT场控：\n'
    '接待官：吴琼\n'
    '摄像官：'
)

# --- B35 官员团队 ---
ws['B35'] = (
    '现任官员团队\n'
    '主席：快乐哥\n'
    '教育副主席：程雪\n'
    '会员副主席：水泽枂\n'
    '公关副主席：水泽枂\n'
    '秘书长：韩羽良\n'
    '财务官：Kim\n'
    '接待官：吴琼'
)

# --- 备稿目标更新 ---
ws['C35'] = ' 1.备稿目标：'
ws['C37'] = ' 2.备稿目标：'
ws['C38'] = ' 3.演讲目标：'

# ==================== 背面 Sheet ====================
ws2 = wb['背面']
ws2['C3'] = '第258次中文线下会议【2026年8月2日】'
# C4 in 背面 has formula =正面!C4, should auto-update

# ==================== 议程反面 Sheet ====================
ws3 = wb['议程反面']
# Update any meeting-specific info here if needed

# ==================== Save ====================
wb.save(dst)
print(f"✓ Excel saved to: {dst}")
print("Done!")
